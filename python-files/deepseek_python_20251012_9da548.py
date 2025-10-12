import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import openpyxl
from openpyxl import Workbook
import os
from datetime import datetime

class MarlinConfiguratorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Marlin 1.1.9 Configurator")
        self.root.geometry("800x600")
        
        # Создаем вкладки
        tab_control = ttk.Notebook(root)
        
        self.tab_basic = ttk.Frame(tab_control)
        self.tab_advanced = ttk.Frame(tab_control)
        self.tab_excel = ttk.Frame(tab_control)
        
        tab_control.add(self.tab_basic, text='Basic')
        tab_control.add(self.tab_advanced, text='Advanced')
        tab_control.add(self.tab_excel, text='Excel')
        
        tab_control.pack(expand=1, fill="both")
        
        # Инициализируем настройки
        self.settings = {
            'profile': tk.StringVar(value="custom"),
            'printer_name': tk.StringVar(value="My Printer"),
            'steps_x': tk.DoubleVar(value=80.0),
            'steps_y': tk.DoubleVar(value=80.0),
            'steps_z': tk.DoubleVar(value=400.0),
            'steps_e': tk.DoubleVar(value=93.0),
            'bed_size_x': tk.IntVar(value=200),
            'bed_size_y': tk.IntVar(value=200),
            'bed_size_z': tk.IntVar(value=200),
            'invert_y': tk.BooleanVar(value=True),
            'thermistor': tk.IntVar(value=1),
            'max_temp_hotend': tk.IntVar(value=275),
            'max_temp_bed': tk.IntVar(value=130),
            'enable_abl': tk.BooleanVar(value=True),
            'enable_bltouch': tk.BooleanVar(value=False),
            'sdcard': tk.BooleanVar(value=True),
            'baudrate': tk.IntVar(value=115200)
        }
        
        # Заполняем вкладки
        self.create_basic_tab()
        self.create_advanced_tab()
        self.create_excel_tab()
        
    def create_basic_tab(self):
        # Выбор профиля
        ttk.Label(self.tab_basic, text="Profile:").grid(column=0, row=0, padx=10, pady=10, sticky=tk.W)
        profile_combo = ttk.Combobox(self.tab_basic, textvariable=self.settings['profile'], 
                                    values=["ender3", "prusa_i3", "custom"])
        profile_combo.grid(column=1, row=0, padx=10, pady=10)
        
        # Имя принтера
        ttk.Label(self.tab_basic, text="Printer Name:").grid(column=0, row=1, padx=10, pady=10, sticky=tk.W)
        ttk.Entry(self.tab_basic, textvariable=self.settings['printer_name']).grid(column=1, row=1, padx=10, pady=10)
        
        # Шаги на мм
        ttk.Label(self.tab_basic, text="Steps per mm:").grid(column=0, row=2, padx=10, pady=10, sticky=tk.W)
        ttk.Label(self.tab_basic, text="X:").grid(column=0, row=3, padx=10, pady=5, sticky=tk.E)
        ttk.Entry(self.tab_basic, textvariable=self.settings['steps_x']).grid(column=1, row=3, padx=10, pady=5)
        ttk.Label(self.tab_basic, text="Y:").grid(column=0, row=4, padx=10, pady=5, sticky=tk.E)
        ttk.Entry(self.tab_basic, textvariable=self.settings['steps_y']).grid(column=1, row=4, padx=10, pady=5)
        ttk.Label(self.tab_basic, text="Z:").grid(column=0, row=5, padx=10, pady=5, sticky=tk.E)
        ttk.Entry(self.tab_basic, textvariable=self.settings['steps_z']).grid(column=1, row=5, padx=10, pady=5)
        ttk.Label(self.tab_basic, text="E:").grid(column=0, row=6, padx=10, pady=5, sticky=tk.E)
        ttk.Entry(self.tab_basic, textvariable=self.settings['steps_e']).grid(column=1, row=6, padx=10, pady=5)
        
        # Размеры стола
        ttk.Label(self.tab_basic, text="Bed Size (mm):").grid(column=0, row=7, padx=10, pady=10, sticky=tk.W)
        ttk.Label(self.tab_basic, text="X:").grid(column=0, row=8, padx=10, pady=5, sticky=tk.E)
        ttk.Entry(self.tab_basic, textvariable=self.settings['bed_size_x']).grid(column=1, row=8, padx=10, pady=5)
        ttk.Label(self.tab_basic, text="Y:").grid(column=0, row=9, padx=10, pady=5, sticky=tk.E)
        ttk.Entry(self.tab_basic, textvariable=self.settings['bed_size_y']).grid(column=1, row=9, padx=10, pady=5)
        ttk.Label(self.tab_basic, text="Z:").grid(column=0, row=10, padx=10, pady=5, sticky=tk.E)
        ttk.Entry(self.tab_basic, textvariable=self.settings['bed_size_z']).grid(column=1, row=10, padx=10, pady=5)
        
        # Дополнительные опции
        ttk.Checkbutton(self.tab_basic, text="Invert Y Axis", variable=self.settings['invert_y']).grid(column=0, row=11, padx=10, pady=5, sticky=tk.W)
        ttk.Checkbutton(self.tab_basic, text="Auto Bed Leveling", variable=self.settings['enable_abl']).grid(column=0, row=12, padx=10, pady=5, sticky=tk.W)
        ttk.Checkbutton(self.tab_basic, text="BLTouch", variable=self.settings['enable_bltouch']).grid(column=0, row=13, padx=10, pady=5, sticky=tk.W)
        ttk.Checkbutton(self.tab_basic, text="SD Card Support", variable=self.settings['sdcard']).grid(column=0, row=14, padx=10, pady=5, sticky=tk.W)
        
        # Кнопка генерации
        ttk.Button(self.tab_basic, text="Generate Configuration", command=self.generate_config).grid(column=0, row=15, padx=10, pady=20)

    def create_advanced_tab(self):
        # Термистор
        ttk.Label(self.tab_advanced, text="Thermistor Type:").grid(column=0, row=0, padx=10, pady=10, sticky=tk.W)
        ttk.Entry(self.tab_advanced, textvariable=self.settings['thermistor']).grid(column=1, row=0, padx=10, pady=10)
        
        # Максимальные температуры
        ttk.Label(self.tab_advanced, text="Max Hotend Temp:").grid(column=0, row=1, padx=10, pady=10, sticky=tk.W)
        ttk.Entry(self.tab_advanced, textvariable=self.settings['max_temp_hotend']).grid(column=1, row=1, padx=10, pady=10)
        ttk.Label(self.tab_advanced, text="Max Bed Temp:").grid(column=0, row=2, padx=10, pady=10, sticky=tk.W)
        ttk.Entry(self.tab_advanced, textvariable=self.settings['max_temp_bed']).grid(column=1, row=2, padx=10, pady=10)
        
        # Скорость порта
        ttk.Label(self.tab_advanced, text="Baudrate:").grid(column=0, row=3, padx=10, pady=10, sticky=tk.W)
        ttk.Entry(self.tab_advanced, textvariable=self.settings['baudrate']).grid(column=1, row=3, padx=10, pady=10)

    def create_excel_tab(self):
        ttk.Button(self.tab_excel, text="Load from Excel", command=self.load_from_excel).grid(column=0, row=0, padx=10, pady=10)
        ttk.Button(self.tab_excel, text="Save to Excel", command=self.save_to_excel).grid(column=1, row=0, padx=10, pady=10)

    def generate_config(self):
        # Получаем настройки из переменных
        profile = self.settings['profile'].get()
        custom_settings = {
            'name': self.settings['printer_name'].get(),
            'steps_per_unit': [
                self.settings['steps_x'].get(),
                self.settings['steps_y'].get(),
                self.settings['steps_z'].get(),
                self.settings['steps_e'].get()
            ],
            'bed_size': [
                self.settings['bed_size_x'].get(),
                self.settings['bed_size_y'].get(),
                self.settings['bed_size_z'].get()
            ],
            'thermistor': self.settings['thermistor'].get(),
            'invert_y': self.settings['invert_y'].get(),
            'max_temp': self.settings['max_temp_hotend'].get(),
            'bed_max_temp': self.settings['max_temp_bed'].get(),
            'enable_abl': self.settings['enable_abl'].get(),
            'enable_bltouch': self.settings['enable_bltouch'].get(),
            'sdcard': self.settings['sdcard'].get(),
            'baudrate': self.settings['baudrate'].get()
        }
        
        # Генерируем конфигурацию
        generator = MarlinConfigGenerator()
        config_content, adv_config_content = generator.generate_configuration(profile, custom_settings)
        
        # Сохраняем в выбранную папку
        output_dir = filedialog.askdirectory(title="Select output directory")
        if output_dir:
            generator.save_configuration(output_dir, config_content, adv_config_content)
            messagebox.showinfo("Success", f"Configuration saved to {output_dir}")

    def load_from_excel(self):
        file_path = filedialog.askopenfilename(title="Open Excel file", filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return
        
        try:
            wb = openpyxl.load_workbook(file_path)
            sheet = wb.active
            
            # Предполагаем, что настройки хранятся в двух колонках: ключ и значение
            for row in sheet.iter_rows(min_row=1, values_only=True):
                if row[0] in self.settings:
                    # В зависимости от типа переменной, преобразуем значение
                    if isinstance(self.settings[row[0]], tk.BooleanVar):
                        self.settings[row[0]].set(bool(row[1]))
                    elif isinstance(self.settings[row[0]], tk.IntVar):
                        self.settings[row[0]].set(int(row[1]))
                    elif isinstance(self.settings[row[0]], tk.DoubleVar):
                        self.settings[row[0]].set(float(row[1]))
                    else:
                        self.settings[row[0]].set(row[1])
            
            messagebox.showinfo("Success", "Settings loaded from Excel")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to load Excel file: {str(e)}")

    def save_to_excel(self):
        file_path = filedialog.asksaveasfilename(title="Save Excel file", defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return
        
        try:
            wb = Workbook()
            sheet = wb.active
            sheet.title = "Marlin Configuration"
            
            # Записываем заголовки
            sheet['A1'] = 'Setting'
            sheet['B1'] = 'Value'
            
            row = 2
            for key, var in self.settings.items():
                sheet[f'A{row}'] = key
                sheet[f'B{row}'] = var.get()
                row += 1
            
            wb.save(file_path)
            messagebox.showinfo("Success", f"Settings saved to {file_path}")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to save Excel file: {str(e)}")

# Класс MarlinConfigGenerator (адаптированный из предыдущего кода)
class MarlinConfigGenerator:
    def __init__(self):
        self.templates = {}
        self.profiles = {}
        self.load_presets()
    
    def load_presets(self):
        """Загрузка предустановленных профилей"""
        self.profiles = {
            'ender3': {
                'name': 'Creality Ender-3',
                'steps_per_unit': [80, 80, 400, 93],
                'bed_size': [220, 220, 250],
                'thermistor': 1,
                'invert_y': True
            },
            'prusa_i3': {
                'name': 'Prusa i3 MK3',
                'steps_per_unit': [100, 100, 400, 280],
                'bed_size': [250, 210, 210],
                'thermistor': 5,
                'invert_y': False
            },
            'custom': {
                'name': 'Custom Printer',
                'steps_per_unit': [100, 100, 400, 93],
                'bed_size': [200, 200, 200],
                'thermistor': 1,
                'invert_y': True
            }
        }
    
    def generate_configuration(self, profile_name, custom_settings=None):
        """Генерация Configuration.h на основе профиля"""
        if profile_name not in self.profiles:
            raise ValueError(f"Профиль {profile_name} не найден")
        
        profile = self.profiles[profile_name].copy()
        if custom_settings:
            profile.update(custom_settings)
        
        config_content = self.build_configuration_h(profile)
        adv_config_content = self.build_configuration_adv_h(profile)
        
        return config_content, adv_config_content
    
    def build_configuration_h(self, settings):
        """Сборка Configuration.h"""
        template = f"""
#ifndef CONFIGURATION_H
#define CONFIGURATION_H

#define CONFIGURATION_H_VERSION 010109

// Generated by Marlin Configurator
// Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
// Profile: {settings['name']}

// ======= BASIC SETTINGS =======
#define SERIAL_PORT 0
#define BAUDRATE {settings.get('baudrate', 115200)}
#define CUSTOM_MACHINE_NAME "{settings['name']}"

// ======= THERMAL SETTINGS =======
#define TEMP_SENSOR_0 {settings['thermistor']}
#define TEMP_SENSOR_BED {settings['thermistor']}

#define HEATER_0_MAXTEMP {settings.get('max_temp', 275)}
#define BED_MAXTEMP {settings.get('bed_max_temp', 130)}

// ======= MECHANICAL SETTINGS =======
#define DEFAULT_AXIS_STEPS_PER_UNIT {{ {settings['steps_per_unit'][0]}, {settings['steps_per_unit'][1]}, {settings['steps_per_unit'][2]}, {settings['steps_per_unit'][3]} }}

#define X_BED_SIZE {settings['bed_size'][0]}
#define Y_BED_SIZE {settings['bed_size'][1]}
#define Z_MAX_POS {settings['bed_size'][2]}

#define INVERT_X_DIR false
#define INVERT_Y_DIR {'true' if settings['invert_y'] else 'false'}
#define INVERT_Z_DIR false

// ======= FEATURES =======
"""
        if settings.get('enable_abl', True):
            template += "#define AUTO_BED_LEVELING_BILINEAR\n"
        
        if settings.get('enable_bltouch', False):
            template += "#define BLTOUCH\n"
        
        if settings.get('sdcard', True):
            template += "#define SDSUPPORT\n"

        template += """
#endif
"""
        return template
    
    def build_configuration_adv_h(self, settings):
        """Сборка Configuration_adv.h"""
        template = f"""
#ifndef CONFIGURATION_ADV_H
#define CONFIGURATION_ADV_H

#define CONFIGURATION_ADV_H_VERSION 010109

// Advanced settings for {settings['name']}

// ======= EXTRUDER SETTINGS =======
#define EXTRUDER_RUNOUT_PREVENT
#define EXTRUDER_RUNOUT_SECONDS 30

// ======= PRECISION SETTINGS =======
#define BABYSTEPPING
#define BABYSTEP_MULTIPLICATOR 1

// ======= SAFETY SETTINGS =======
#define THERMAL_PROTECTION_HOTENDS
#define THERMAL_PROTECTION_BED

#endif
"""
        return template
    
    def save_configuration(self, output_dir, config_content, adv_config_content):
        """Сохранение конфигурационных файлов"""
        os.makedirs(output_dir, exist_ok=True)
        
        with open(os.path.join(output_dir, 'Configuration.h'), 'w') as f:
            f.write(config_content)
        
        with open(os.path.join(output_dir, 'Configuration_adv.h'), 'w') as f:
            f.write(adv_config_content)
        
        print(f"Конфигурация сохранена в {output_dir}")

if __name__ == "__main__":
    root = tk.Tk()
    app = MarlinConfiguratorApp(root)
    root.mainloop()